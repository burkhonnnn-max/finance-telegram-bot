# -*- coding: utf-8 -*-
"""
Excel hisobot yaratuvchi modul (openpyxl orqali)
Oylik va davriy moliyaviy hisobotlarni zamonaviy, rangli va qulay Excel (.xlsx) fayliga aylantiradi.
"""

import io
from datetime import datetime
from typing import List, Dict, Any

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from locales import localize_category


MONTHS_UZ = {
    1: "Yanvar", 2: "Fevral", 3: "Mart", 4: "Aprel",
    5: "May", 6: "Iyun", 7: "Iyul", 8: "Avgust",
    9: "Sentyabr", 10: "Oktyabr", 11: "Noyabr", 12: "Dekabr"
}

MONTHS_RU = {
    1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель",
    5: "Май", 6: "Июнь", 7: "Июль", 8: "Август",
    9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь"
}


def create_monthly_excel_report(
    user_name: str,
    user_lang: str,
    year: int,
    month: int,
    stats: Dict[str, Any],
    transactions: List[Dict[str, Any]]
) -> bytes:
    """
    Foydalanuvchining oylik hisobotini chiroyli va professional Excel (.xlsx) formatida yaratish.
    Natija sifatida bytes qaytaradi.
    """
    wb = openpyxl.Workbook()
    
    # Tillarga mos matnlar
    is_ru = (user_lang == "ru")
    month_name = MONTHS_RU.get(month, "") if is_ru else MONTHS_UZ.get(month, "")
    
    # ------------------ STYLES ------------------
    font_family = "Segoe UI"
    
    title_font = Font(name=font_family, size=15, bold=True, color="FFFFFF")
    section_font = Font(name=font_family, size=12, bold=True, color="1F4E79")
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    bold_font = Font(name=font_family, size=10, bold=True)
    regular_font = Font(name=font_family, size=10)
    meta_font = Font(name=font_family, size=9, italic=True, color="555555")
    
    income_font = Font(name=font_family, size=10, bold=True, color="1B5E20")
    expense_font = Font(name=font_family, size=10, bold=True, color="B71C1C")
    
    primary_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    accent_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    sub_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    
    income_row_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    expense_row_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
    summary_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="D0D7DE")
    border_box = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    thick_bottom_side = Side(border_style="medium", color="1F4E79")
    total_top_side = Side(border_style="thin", color="1F4E79")
    total_bottom_side = Side(border_style="double", color="1F4E79")
    total_border = Border(top=total_top_side, bottom=total_bottom_side, left=thin_border_side, right=thin_border_side)
    
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    
    number_format = "#,##0"
    currency_format = '#,##0" so\'m"' if not is_ru else '#,##0" сум"'
    
    # ==========================================================
    # 1-SHEET: UMUMIY HISOBOT VA XULOSA (Summary)
    # ==========================================================
    ws_summary = wb.active
    ws_summary.title = "Xulosa" if not is_ru else "Сводка"
    ws_summary.views.sheetView[0].showGridLines = True
    
    # 1. Title Banner
    ws_summary.merge_cells("A1:E1")
    cell_title = ws_summary["A1"]
    if is_ru:
        cell_title.value = f"📅 ФИНАНСОВЫЙ ОТЧЁТ ЗА {month_name.upper()} {year} ГОДА"
    else:
        cell_title.value = f"📅 {year}-YIL {month_name.upper()} OYI UCHUN MOLIYAVIY HISOBOT"
    cell_title.font = title_font
    cell_title.fill = primary_fill
    cell_title.alignment = center_align
    ws_summary.row_dimensions[1].height = 40
    
    # Meta ma'lumotlar
    ws_summary.merge_cells("A2:E2")
    cell_meta = ws_summary["A2"]
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    if is_ru:
        cell_meta.value = f"Пользователь: {user_name}  |  Сформировано: {now_str}  |  @ybr_money_bot"
    else:
        cell_meta.value = f"Foydalanuvchi: {user_name}  |  Yaratilgan vaqt: {now_str}  |  @ybr_money_bot"
    cell_meta.font = meta_font
    cell_meta.fill = sub_fill
    cell_meta.alignment = center_align
    ws_summary.row_dimensions[2].height = 20
    
    # 2. Asosiy Moliyaviy Ko'rsatkichlar (KPIs)
    row = 4
    ws_summary.cell(row=row, column=1, value="📊 ASOSIY KO'RSATKICHLAR" if not is_ru else "📊 ОСНОВНЫЕ ПОКАЗАТЕЛИ").font = section_font
    row += 1
    
    kpi_headers = ["Ko'rsatkich", "Qiymat"] if not is_ru else ["Показатель", "Значение"]
    for col_idx, h in enumerate(kpi_headers, start=1):
        c = ws_summary.cell(row=row, column=col_idx, value=h)
        c.font = header_font
        c.fill = accent_fill
        c.alignment = left_align if col_idx == 1 else right_align
        c.border = border_box
    ws_summary.row_dimensions[row].height = 24
    row += 1
    
    total_income = stats.get("income", 0.0)
    total_expense = stats.get("expense", 0.0)
    diff = stats.get("difference", 0.0)
    total_ops = len(transactions)
    
    kpi_items = [
        ("🟢 " + ("Jami kirim (daromad)" if not is_ru else "Всего доходов"), total_income, income_font, income_row_fill),
        ("🔴 " + ("Jami chiqim (xarajat)" if not is_ru else "Всего расходов"), total_expense, expense_font, expense_row_fill),
        ("💰 " + ("Sof jamg'arma (qoldiq)" if not is_ru else "Чистый остаток (накопления)"), diff, bold_font, summary_fill),
        ("📝 " + ("Jami amallar soni" if not is_ru else "Всего операций"), total_ops, regular_font, None),
    ]
    
    for label, val, f_style, b_fill in kpi_items:
        c1 = ws_summary.cell(row=row, column=1, value=label)
        c1.font = f_style
        c1.border = border_box
        if b_fill:
            c1.fill = b_fill
            
        c2 = ws_summary.cell(row=row, column=2, value=val)
        c2.font = f_style
        c2.border = border_box
        c2.alignment = right_align
        if isinstance(val, (int, float)) and label.startswith(("🟢", "🔴", "💰")):
            c2.number_format = currency_format
        else:
            c2.number_format = number_format
        if b_fill:
            c2.fill = b_fill
            
        ws_summary.row_dimensions[row].height = 22
        row += 1
        
    row += 1
    
    # 3. Chiqimlar toifalar bo'yicha (Expenses breakdown)
    ws_summary.cell(row=row, column=1, value="🛒 CHIQIMLAR TOIFALAR BO'YICHA" if not is_ru else "🛒 РАСХОДЫ ПО КАТЕГОРИЯМ").font = section_font
    row += 1
    
    exp_headers = ["№", "Toifa nomi", "Jami summa", "Ulush (Foiz)", "Amallar soni"] if not is_ru else ["№", "Категория", "Сумма", "Доля (%)", "Кол-во"]
    for col_idx, h in enumerate(exp_headers, start=1):
        c = ws_summary.cell(row=row, column=col_idx, value=h)
        c.font = header_font
        c.fill = accent_fill
        c.alignment = center_align if col_idx in [1, 4, 5] else (left_align if col_idx == 2 else right_align)
        c.border = border_box
    ws_summary.row_dimensions[row].height = 24
    row += 1
    
    categories_expense = stats.get("categories_expense", [])
    if categories_expense:
        for idx, cat in enumerate(categories_expense, start=1):
            c_num = ws_summary.cell(row=row, column=1, value=idx)
            c_name = ws_summary.cell(row=row, column=2, value=localize_category(cat["category"], user_lang))
            c_tot = ws_summary.cell(row=row, column=3, value=cat["total"])
            c_pct = ws_summary.cell(row=row, column=4, value=f"{cat['percentage']}%")
            c_cnt = ws_summary.cell(row=row, column=5, value=cat["count"])
            
            c_num.alignment = center_align
            c_name.alignment = left_align
            c_tot.alignment = right_align
            c_tot.number_format = currency_format
            c_pct.alignment = center_align
            c_cnt.alignment = center_align
            c_cnt.number_format = number_format
            
            for cell in [c_num, c_name, c_tot, c_pct, c_cnt]:
                cell.font = regular_font
                cell.border = border_box
                
            ws_summary.row_dimensions[row].height = 20
            row += 1
    else:
        ws_summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        c_empty = ws_summary.cell(row=row, column=1, value="Bu oyda hech qanday chiqim bo'lmagan" if not is_ru else "В этом месяце расходов не было")
        c_empty.font = meta_font
        c_empty.alignment = center_align
        row += 1
        
    row += 1
    
    # 4. Kirimlar toifalar bo'yicha (Income breakdown)
    ws_summary.cell(row=row, column=1, value="💼 KIRIMLAR TOIFALAR BO'YICHA" if not is_ru else "💼 ДОХОДЫ ПО КАТЕГОРИЯМ").font = section_font
    row += 1
    
    inc_headers = ["№", "Toifa nomi", "Jami summa", "Ulush (Foiz)", "Amallar soni"] if not is_ru else ["№", "Категория", "Сумма", "Доля (%)", "Кол-во"]
    for col_idx, h in enumerate(inc_headers, start=1):
        c = ws_summary.cell(row=row, column=col_idx, value=h)
        c.font = header_font
        c.fill = accent_fill
        c.alignment = center_align if col_idx in [1, 4, 5] else (left_align if col_idx == 2 else right_align)
        c.border = border_box
    ws_summary.row_dimensions[row].height = 24
    row += 1
    
    categories_income = stats.get("categories_income", [])
    if categories_income:
        for idx, cat in enumerate(categories_income, start=1):
            c_num = ws_summary.cell(row=row, column=1, value=idx)
            c_name = ws_summary.cell(row=row, column=2, value=localize_category(cat["category"], user_lang))
            c_tot = ws_summary.cell(row=row, column=3, value=cat["total"])
            c_pct = ws_summary.cell(row=row, column=4, value=f"{cat['percentage']}%")
            c_cnt = ws_summary.cell(row=row, column=5, value=cat["count"])
            
            c_num.alignment = center_align
            c_name.alignment = left_align
            c_tot.alignment = right_align
            c_tot.number_format = currency_format
            c_pct.alignment = center_align
            c_cnt.alignment = center_align
            c_cnt.number_format = number_format
            
            for cell in [c_num, c_name, c_tot, c_pct, c_cnt]:
                cell.font = regular_font
                cell.border = border_box
                
            ws_summary.row_dimensions[row].height = 20
            row += 1
    else:
        ws_summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        c_empty = ws_summary.cell(row=row, column=1, value="Bu oyda hech qanday kirim bo'lmagan" if not is_ru else "В этом месяце доходов не было")
        c_empty.font = meta_font
        c_empty.alignment = center_align
        row += 1

    # Auto-width for summary sheet
    for col in ws_summary.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Sarlavhalarni hisobga olmaslik
            if cell.row in [1, 2]:
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 14)
        
    # ==========================================================
    # 2-SHEET: BARCHA TRANZAKSIYALAR RO'YXATI (Detailed Transactions)
    # ==========================================================
    ws_trans = wb.create_sheet(title="Amallar ro'yxati" if not is_ru else "Все операции")
    ws_trans.views.sheetView[0].showGridLines = True
    
    # Headers
    tx_headers = (
        ["№", "Sana va vaqt", "Turi", "Summa", "Toifa", "Izoh"]
        if not is_ru else
        ["№", "Дата и время", "Тип", "Сумма", "Категория", "Примечание"]
    )
    
    for col_idx, h in enumerate(tx_headers, start=1):
        c = ws_trans.cell(row=1, column=col_idx, value=h)
        c.font = header_font
        c.fill = primary_fill
        c.alignment = center_align if col_idx in [1, 2, 3] else (right_align if col_idx == 4 else left_align)
        c.border = border_box
    ws_trans.row_dimensions[1].height = 28
    
    # Freeze top row
    ws_trans.freeze_panes = "A2"
    
    tx_row = 2
    for idx, tx in enumerate(transactions, start=1):
        is_income = (tx.get("type") == "income")
        type_label = ("➕ Kirim" if is_income else "➖ Chiqim") if not is_ru else ("➕ Доход" if is_income else "➖ Расход")
        
        c_id = ws_trans.cell(row=tx_row, column=1, value=idx)
        c_date = ws_trans.cell(row=tx_row, column=2, value=tx.get("created_at", ""))
        c_type = ws_trans.cell(row=tx_row, column=3, value=type_label)
        c_amt = ws_trans.cell(row=tx_row, column=4, value=tx.get("amount", 0.0))
        c_cat = ws_trans.cell(row=tx_row, column=5, value=localize_category(tx.get("category", ""), user_lang))
        c_com = ws_trans.cell(row=tx_row, column=6, value=tx.get("comment") or "-")
        
        c_id.alignment = center_align
        c_date.alignment = center_align
        c_type.alignment = center_align
        c_amt.alignment = right_align
        c_amt.number_format = currency_format
        c_cat.alignment = left_align
        c_com.alignment = left_align
        
        row_fill = income_row_fill if is_income else expense_row_fill
        type_font = income_font if is_income else expense_font
        
        for c in [c_id, c_date, c_cat, c_com]:
            c.font = regular_font
            c.border = border_box
            c.fill = row_fill
            
        c_type.font = type_font
        c_type.border = border_box
        c_type.fill = row_fill
        
        c_amt.font = type_font
        c_amt.border = border_box
        c_amt.fill = row_fill
        
        ws_trans.row_dimensions[tx_row].height = 20
        tx_row += 1
        
    # Auto-filter on transactions table
    if transactions:
        ws_trans.auto_filter.ref = f"A1:F{tx_row - 1}"
        
    # Auto-width for transactions sheet
    for col in ws_trans.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_trans.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # Bytes ga o'girish
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
